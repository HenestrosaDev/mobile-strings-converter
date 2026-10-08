import io

import openpyxl

from ..model import Catalog
from . import table


def parse(data: bytes, locale: str, with_comments: bool = False) -> Catalog:
    """Reads the active sheet. See `table` for the supported columns."""

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("The file provided has no sheets.")
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    return table.from_table(rows[0] if rows else None, rows[1:], locale)


def serialize(catalog: Catalog) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.worksheets[0]

    for row in table.to_table(catalog):
        sheet.append(row)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()

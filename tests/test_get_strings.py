import unittest
from pathlib import Path
from tempfile import NamedTemporaryFile, TemporaryDirectory
from unittest import mock

import openpyxl
from base_tests import BaseTests

from mobile_strings_converter import ConversionWarning, parse
from mobile_strings_converter.converter import get_strings


class TestGetStringsIos(BaseTests.GetStringsTest):
    def setUp(self):
        self.data = """
        //      "chinese" = "  欢迎来到我的申请  "   ;
            //"escaped_quote"="MyA\\\"pp";
        "hindi" = "मेरे ऐप का आनंद लें";
        "korean" = "내 앱을 즐기세요";
        """
        self.extension = ".strings"


class TestGetStringsAndroid(BaseTests.GetStringsTest):
    def setUp(self):
        self.data = """
        <?xml version="1.0" encoding="UTF-8"?>
        <resources>
            <!--<string name="chinese">  欢迎来到我的申请  </string>-->
            <!--    <string name="escaped_quote">MyA\\\"pp</string>  -->
            <string name="hindi">मेरे ऐप का आनंद लें</string>
                    <string name="korean">내 앱을 즐기세요</string>
            "3" = "3";
        </resources>
        """
        self.extension = ".xml"


class TestAndroidParsing(unittest.TestCase):
    def _get_strings(
        self, data: str, with_comments: bool = False
    ) -> list[tuple[str, str]]:
        with NamedTemporaryFile(
            suffix=".xml", mode="w", delete=False, encoding="utf-8"
        ) as f:
            f.write(data)
            filepath = Path(f.name)

        try:
            return get_strings(filepath, with_comments)
        finally:
            filepath.unlink()

    def test_attributes_multiline_plurals_and_arrays(self):
        data = """<?xml version="1.0" encoding="utf-8"?>
        <resources>
            <string name="app_name" translatable="false">My App</string>
            <string
                name="multiline">first\\nsecond</string>
            <string name="quoted">"It's quoted"</string>
            <plurals name="songs">
                <item quantity="one">%d song</item>
                <item quantity="other">%d songs</item>
            </plurals>
            <string-array name="planets">
                <item>Mercury</item>
            </string-array>
            <!-- Just a regular comment -->
        </resources>
        """

        self.assertEqual(
            [
                ("app_name", "My App"),
                ("multiline", "first\nsecond"),
                ("quoted", "It's quoted"),
                ("songs[one]", "%d song"),
                ("songs[other]", "%d songs"),
                ("planets[0]", "Mercury"),
            ],
            self._get_strings(data, with_comments=True),
        )


class TestIosParsing(unittest.TestCase):
    def _get_strings(
        self, data: str, with_comments: bool = False
    ) -> list[tuple[str, str]]:
        with NamedTemporaryFile(
            suffix=".strings", mode="w", delete=False, encoding="utf-8"
        ) as f:
            f.write(data)
            filepath = Path(f.name)

        try:
            return get_strings(filepath, with_comments)
        finally:
            filepath.unlink()

    def test_block_comments(self):
        data = """
        /* Title of the main screen */
        "title" = "Home";
        /* "old_title" = "Start"; */
        "url" = "https://example.com"; // Trailing comment
        "unicode" = "\\U00e9t\\u00e9";
        """

        self.assertEqual(
            [("title", "Home"), ("url", "https://example.com"), ("unicode", "été")],
            self._get_strings(data, with_comments=False),
        )
        self.assertEqual(
            [
                ("title", "Home"),
                ("old_title", "Start"),
                ("url", "https://example.com"),
                ("unicode", "été"),
            ],
            self._get_strings(data, with_comments=True),
        )


class TestReaders(unittest.TestCase):
    def _get_strings(self, data: str, extension: str) -> list[tuple[str, str]]:
        with NamedTemporaryFile(
            suffix=extension, mode="w", delete=False, encoding="utf-8"
        ) as f:
            f.write(data)
            filepath = Path(f.name)

        try:
            return get_strings(filepath)
        finally:
            filepath.unlink()

    def test_csv_with_bom_blank_rows_and_extra_columns(self):
        data = "\ufeffname,value\nhello,Hello,extra\n\n,no name\nempty\n"

        self.assertEqual(
            [("hello", "Hello"), ("empty", "")], self._get_strings(data, ".csv")
        )

    def test_json_object(self):
        data = '{"hello": "Hello", "count": 3}'

        self.assertEqual(
            [("hello", "Hello"), ("count", "3")], self._get_strings(data, ".json")
        )

    def test_yaml_with_non_string_values(self):
        data = "hello: Hello\ncount: 3\nenabled: yes\nempty:\n"

        self.assertEqual(
            [("hello", "Hello"), ("count", "3"), ("enabled", "True"), ("empty", "")],
            self._get_strings(data, ".yaml"),
        )

    def test_empty_yaml(self):
        self.assertEqual([], self._get_strings("", ".yaml"))

    def test_xlsx_with_numbers_and_blank_rows(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.worksheets[0]
        for row in [("NAME", "VALUE"), ("count", 3), (None, None), ("empty", None)]:
            sheet.append(list(row))

        with TemporaryDirectory() as temp_dir:
            filepath = Path(temp_dir) / "strings.xlsx"
            workbook.save(filepath)

            self.assertEqual([("count", "3"), ("empty", "")], get_strings(filepath))

    def test_xlsx_without_sheets(self):
        with (
            mock.patch("openpyxl.load_workbook") as load_workbook,
            self.assertRaisesRegex(ValueError, "no sheets"),
        ):
            load_workbook.return_value.active = None
            parse(b"", ".xlsx")

    def test_unsupported_android_resources(self):
        data = (
            b'<resources><integer-array name="numbers"><item>1</item></integer-array>'
            b'<string name="hello">Hello</string></resources>'
        )

        with self.assertWarnsRegex(ConversionWarning, "integer-array"):
            catalog = parse(data, ".xml")

        self.assertEqual([("hello", "Hello")], catalog.to_pairs())

    def test_uppercase_extension(self):
        self.assertEqual(
            [("hello", "Hello")], self._get_strings('"hello" = "Hello";', ".STRINGS")
        )


if __name__ == "__main__":
    unittest.main()
